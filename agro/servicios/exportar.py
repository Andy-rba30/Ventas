"""Exportación a Excel con openpyxl (sin pandas)."""
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNAS_BOLETAS = ["Fecha", "Hora", "Tipo", "Cliente/Proveedor", "Encargada", "Total (S/.)", "Estado", "Ref"]
COLUMNAS_LINEAS = ["Fecha", "Tipo", "Cliente/Proveedor", "Producto", "Cantidad", "P. Unit (S/.)", "Subtotal (S/.)", "Ref boleta"]


def escribir_hoja(ws, columnas, filas):
    """Cabecera en negrita, filas de datos y ancho de columna según el contenido (mín 8, máx 45)."""
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append(list(fila))
    for i, col in enumerate(columnas, start=1):
        ancho = len(str(col))
        for fila in filas[:500]:
            ancho = max(ancho, len(str(fila[i - 1])))
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 8), 45)


def boletas_a_excel(ruta, boletas):
    """Dos hojas: Boletas (una fila por boleta o pago) y Lineas (una fila por producto)."""
    filas_boletas = [(b.fecha, b.hora, b.tipo, b.persona, b.encargada, b.total, b.estado, b.clave) for b in boletas]
    filas_lineas = [(b.fecha, b.tipo, b.persona, l.producto, l.cantidad, l.precio_unit, l.total, b.clave)
                    for b in boletas for l in b.lineas]
    wb = Workbook()
    ws = wb.active
    ws.title = "Boletas"
    escribir_hoja(ws, COLUMNAS_BOLETAS, filas_boletas)
    escribir_hoja(wb.create_sheet("Lineas"), COLUMNAS_LINEAS, filas_lineas)
    wb.save(ruta)
    return ruta
